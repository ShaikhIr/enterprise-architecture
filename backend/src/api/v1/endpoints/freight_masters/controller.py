"""
Freight Masters API Controller — CRUD for Air, Sea, Sea CBM, Vehicle Type, Local masters.
Includes rate sub-endpoints and Excel import/template downloads.
"""

import io
from datetime import datetime, timezone
from decimal import Decimal
from uuid import UUID

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlalchemy import select, text as _text
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from src.api.v1.dependencies import get_current_active_user
from src.domain.entities.user import User
from src.infrastructure.database.session import get_db_session
from src.infrastructure.database.models.masters.country_model import CountryMasterModel
from src.infrastructure.database.models.masters.city_model import CityMasterModel
from src.infrastructure.database.models.masters.air_master_model import (
    AirMasterModel,
    AirMasterRateModel,
)
from src.infrastructure.database.models.masters.sea_master_model import (
    SeaMasterModel,
    SeaMasterRateModel,
)
from src.infrastructure.database.models.masters.sea_cbm_master_model import SeaCbmMasterModel
from src.infrastructure.database.models.masters.vehicle_type_master_model import VehicleTypeMasterModel
from src.infrastructure.database.models.masters.local_master_model import (
    LocalMasterModel,
    LocalMasterRateModel,
)
from src.api.v1.endpoints.freight_masters.schemas import (
    CountryCreateRequest, CountryUpdateRequest, CountryResponse,
    CityCreateRequest, CityUpdateRequest, CityResponse,
    AirMasterCreateRequest, AirMasterUpdateRequest, AirMasterResponse,
    SeaMasterCreateRequest, SeaMasterUpdateRequest, SeaMasterResponse,
    SeaCbmCreateRequest, SeaCbmUpdateRequest, SeaCbmResponse,
    VehicleTypeCreateRequest, VehicleTypeUpdateRequest, VehicleTypeResponse,
    LocalMasterCreateRequest, LocalMasterUpdateRequest, LocalMasterResponse,
    RateCreateRequest, RateResponse,
)

router = APIRouter(prefix="/freight-masters", tags=["Freight Masters"])


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def _air_to_resp(record) -> AirMasterResponse:
    """Map AirMasterModel ORM object to response schema."""
    return AirMasterResponse(
        id=record.id,
        to_country_id=record.to_country_id,
        country_code=record.to_country.country_code if record.to_country else None,
        country_name=record.to_country.country_name if record.to_country else None,
        product_type=record.product_type,
        min_slab=record.min_slab,
        max_slab=record.max_slab,
        slab_name=record.slab_name,
        rates=[
            RateResponse(
                id=r.id,
                rate=float(r.rate),
                valid_from=r.valid_from.isoformat() if r.valid_from else None,
                valid_till=r.valid_till.isoformat() if r.valid_till else None,
                is_active=r.is_active,
            )
            for r in (record.rates or [])
        ],
        is_active=record.is_active,
    )


def _sea_to_resp(record) -> SeaMasterResponse:
    """Map SeaMasterModel ORM object to response schema."""
    return SeaMasterResponse(
        id=record.id,
        to_country_id=record.to_country_id,
        country_code=record.to_country.country_code if record.to_country else None,
        country_name=record.to_country.country_name if record.to_country else None,
        product_type=record.product_type,
        slab_name=record.slab_name,
        currency=record.currency,
        rates=[
            RateResponse(
                id=r.id,
                rate=float(r.rate),
                valid_from=r.valid_from.isoformat() if r.valid_from else None,
                valid_till=r.valid_till.isoformat() if r.valid_till else None,
                is_active=r.is_active,
            )
            for r in (record.rates or [])
        ],
        is_active=record.is_active,
    )


def _local_to_resp(record) -> LocalMasterResponse:
    """Map LocalMasterModel ORM object to response schema."""
    return LocalMasterResponse(
        id=record.id,
        from_city_id=record.from_city_id,
        from_city_name=record.from_city.city_name if record.from_city else None,
        to_city_id=record.to_city_id,
        to_city_name=record.to_city.city_name if record.to_city else None,
        product_type=record.product_type,
        vehicle_type_id=record.vehicle_type_id,
        vehicle_type_name=record.vehicle_type_ref.vehicle_type if record.vehicle_type_ref else None,
        min_slab=record.min_slab,
        max_slab=record.max_slab,
        rates=[
            RateResponse(
                id=r.id,
                rate=float(r.rate),
                valid_from=r.valid_from.isoformat() if r.valid_from else None,
                valid_till=r.valid_till.isoformat() if r.valid_till else None,
                is_active=r.is_active,
            )
            for r in (record.rates or [])
        ],
        is_active=record.is_active,
    )


# ============================================================
# COUNTRY MASTER — CRUD
# ============================================================

@router.get("/countries", response_model=list[CountryResponse])
async def list_countries(
    include_inactive: bool = False,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """List countries. By default only active records are returned."""
    query = select(CountryMasterModel).order_by(CountryMasterModel.country_name)
    if not include_inactive:
        query = query.where(CountryMasterModel.is_active == True)  # noqa: E712
    result = await session.execute(query)
    return [CountryResponse.model_validate(r) for r in result.scalars().all()]


@router.post("/countries", response_model=CountryResponse, status_code=status.HTTP_201_CREATED)
async def create_country(
    request: CountryCreateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Create a new country. Duplicate country_code not allowed."""
    existing = await session.execute(
        select(CountryMasterModel).where(
            CountryMasterModel.country_code.ilike(request.country_code)
        )
    )
    if existing.scalar_one_or_none():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Country code already exists.")
    record = CountryMasterModel(
        country_name=request.country_name,
        country_code=request.country_code.upper(),
        is_active=True,
        created_by=current_user.username,
        modified_by=current_user.username,
    )
    session.add(record)
    await session.flush()
    await session.refresh(record)
    return CountryResponse.model_validate(record)


@router.put("/countries/{record_id}", response_model=CountryResponse)
async def update_country(
    record_id: UUID,
    request: CountryUpdateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Update a country."""
    result = await session.execute(
        select(CountryMasterModel).where(CountryMasterModel.id == str(record_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Country not found")
    if request.country_code is not None:
        dup = await session.execute(
            select(CountryMasterModel).where(
                CountryMasterModel.country_code.ilike(request.country_code),
                CountryMasterModel.id != str(record_id),
            )
        )
        if dup.scalar_one_or_none():
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Country code already exists.")
        record.country_code = request.country_code.upper()
    if request.country_name is not None:
        record.country_name = request.country_name
    if request.is_active is not None:
        record.is_active = request.is_active
    record.modified_by = current_user.username
    await session.flush()
    await session.refresh(record)
    return CountryResponse.model_validate(record)


@router.delete("/countries/{record_id}", status_code=status.HTTP_200_OK)
async def delete_country(
    record_id: UUID,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Soft-delete a country."""
    result = await session.execute(
        select(CountryMasterModel).where(CountryMasterModel.id == str(record_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Country not found")
    record.is_active = False
    record.modified_by = current_user.username
    await session.flush()
    return {"message": "Country deactivated successfully"}


# ============================================================
# CITY MASTER — CRUD
# ============================================================

@router.get("/cities", response_model=list[CityResponse])
async def list_cities(
    include_inactive: bool = False,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """List cities. By default only active records are returned."""
    query = select(CityMasterModel).order_by(CityMasterModel.city_name)
    if not include_inactive:
        query = query.where(CityMasterModel.is_active == True)  # noqa: E712
    result = await session.execute(query)
    return [CityResponse.model_validate(r) for r in result.scalars().all()]


@router.post("/cities", response_model=CityResponse, status_code=status.HTTP_201_CREATED)
async def create_city(
    request: CityCreateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Create a new city."""
    record = CityMasterModel(
        city_name=request.city_name,
        country_id=str(request.country_id) if request.country_id else None,
        is_active=True,
        created_by=current_user.username,
        modified_by=current_user.username,
    )
    session.add(record)
    await session.flush()
    await session.refresh(record)
    return CityResponse.model_validate(record)


@router.put("/cities/{record_id}", response_model=CityResponse)
async def update_city(
    record_id: UUID,
    request: CityUpdateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Update a city."""
    result = await session.execute(
        select(CityMasterModel).where(CityMasterModel.id == str(record_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="City not found")
    if request.city_name is not None:
        record.city_name = request.city_name
    if request.country_id is not None:
        record.country_id = str(request.country_id)
    if request.is_active is not None:
        record.is_active = request.is_active
    record.modified_by = current_user.username
    await session.flush()
    await session.refresh(record)
    return CityResponse.model_validate(record)


@router.delete("/cities/{record_id}", status_code=status.HTTP_200_OK)
async def delete_city(
    record_id: UUID,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Soft-delete a city."""
    result = await session.execute(
        select(CityMasterModel).where(CityMasterModel.id == str(record_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="City not found")
    record.is_active = False
    record.modified_by = current_user.username
    await session.flush()
    return {"message": "City deactivated successfully"}


# ============================================================
# AIR MASTER (Freight Slab) — CRUD
# ============================================================

@router.get("/air-masters", response_model=list[AirMasterResponse])
async def list_air_masters(
    include_inactive: bool = False,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """List Air Masters with rates. By default only active records."""
    query = (
        select(AirMasterModel)
        .options(selectinload(AirMasterModel.to_country), selectinload(AirMasterModel.rates))
        .order_by(AirMasterModel.created_date.desc())
    )
    if not include_inactive:
        query = query.where(AirMasterModel.is_active == True)  # noqa: E712
    result = await session.execute(query)
    return [_air_to_resp(r) for r in result.scalars().all()]


@router.post("/air-masters", response_model=AirMasterResponse, status_code=status.HTTP_201_CREATED)
async def create_air_master(
    request: AirMasterCreateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Create a new Air Master record."""
    slab_name = f"{request.min_slab}-{request.max_slab} KG"
    record = AirMasterModel(
        to_country_id=str(request.to_country_id),
        product_type=request.product_type,
        min_slab=request.min_slab,
        max_slab=request.max_slab,
        slab_name=slab_name,
        is_active=True,
        created_by=current_user.username,
        modified_by=current_user.username,
    )
    session.add(record)
    await session.flush()
    res = await session.execute(
        select(AirMasterModel)
        .options(selectinload(AirMasterModel.to_country), selectinload(AirMasterModel.rates))
        .where(AirMasterModel.id == record.id)
    )
    return _air_to_resp(res.scalar_one())


@router.put("/air-masters/{record_id}", response_model=AirMasterResponse)
async def update_air_master(
    record_id: UUID,
    request: AirMasterUpdateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Update an Air Master record."""
    result = await session.execute(
        select(AirMasterModel).where(AirMasterModel.id == str(record_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Air Master not found")
    if request.to_country_id is not None:
        record.to_country_id = str(request.to_country_id)
    if request.product_type is not None:
        record.product_type = request.product_type
    if request.min_slab is not None:
        record.min_slab = request.min_slab
    if request.max_slab is not None:
        record.max_slab = request.max_slab
    if request.min_slab is not None or request.max_slab is not None:
        record.slab_name = f"{record.min_slab}-{record.max_slab} KG"
    if request.is_active is not None:
        record.is_active = request.is_active
    record.modified_by = current_user.username
    await session.flush()
    res = await session.execute(
        select(AirMasterModel)
        .options(selectinload(AirMasterModel.to_country), selectinload(AirMasterModel.rates))
        .where(AirMasterModel.id == str(record_id))
    )
    return _air_to_resp(res.scalar_one())


@router.delete("/air-masters/{record_id}", status_code=status.HTTP_200_OK)
async def delete_air_master(
    record_id: UUID,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Soft-delete an Air Master record."""
    result = await session.execute(
        select(AirMasterModel).where(AirMasterModel.id == str(record_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Air Master not found")
    record.is_active = False
    record.modified_by = current_user.username
    await session.flush()
    return {"message": "Air Master deactivated successfully"}


# --- Air Master Rate sub-endpoints ---

@router.post("/air-masters/{master_id}/rates", response_model=RateResponse, status_code=status.HTTP_201_CREATED)
async def add_air_master_rate(
    master_id: UUID,
    request: RateCreateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Add a rate to an Air Master record."""
    check = await session.execute(
        select(AirMasterModel).where(AirMasterModel.id == str(master_id))
    )
    if not check.scalar_one_or_none():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Air Master not found")
    vf = datetime.fromisoformat(request.valid_from.replace("Z", "+00:00"))
    vt = datetime.fromisoformat(request.valid_till.replace("Z", "+00:00"))
    record = AirMasterRateModel(
        air_master_id=str(master_id),
        rate=Decimal(str(request.rate)),
        valid_from=vf,
        valid_till=vt,
        is_active=True,
        created_by=current_user.username,
        modified_by=current_user.username,
    )
    session.add(record)
    await session.flush()
    return RateResponse(
        id=record.id,
        rate=float(record.rate),
        valid_from=record.valid_from.isoformat(),
        valid_till=record.valid_till.isoformat(),
        is_active=record.is_active,
    )


@router.delete("/air-masters/rates/{rate_id}", status_code=status.HTTP_200_OK)
async def delete_air_master_rate(
    rate_id: UUID,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Soft-delete an Air Master rate."""
    result = await session.execute(
        select(AirMasterRateModel).where(AirMasterRateModel.id == str(rate_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Rate not found")
    record.is_active = False
    record.modified_by = current_user.username
    await session.flush()
    return {"message": "Rate deactivated successfully"}


# ============================================================
# SEA MASTER — CRUD
# ============================================================

@router.get("/sea-masters", response_model=list[SeaMasterResponse])
async def list_sea_masters(
    include_inactive: bool = False,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """List Sea Masters with rates."""
    query = (
        select(SeaMasterModel)
        .options(selectinload(SeaMasterModel.to_country), selectinload(SeaMasterModel.rates))
        .order_by(SeaMasterModel.created_date.desc())
    )
    if not include_inactive:
        query = query.where(SeaMasterModel.is_active == True)  # noqa: E712
    result = await session.execute(query)
    return [_sea_to_resp(r) for r in result.scalars().all()]


@router.post("/sea-masters", response_model=SeaMasterResponse, status_code=status.HTTP_201_CREATED)
async def create_sea_master(
    request: SeaMasterCreateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Create a new Sea Master record."""
    record = SeaMasterModel(
        to_country_id=str(request.to_country_id),
        product_type=request.product_type,
        slab_name=request.slab_name,
        currency=request.currency,
        is_active=True,
        created_by=current_user.username,
        modified_by=current_user.username,
    )
    session.add(record)
    await session.flush()
    res = await session.execute(
        select(SeaMasterModel)
        .options(selectinload(SeaMasterModel.to_country), selectinload(SeaMasterModel.rates))
        .where(SeaMasterModel.id == record.id)
    )
    return _sea_to_resp(res.scalar_one())


@router.put("/sea-masters/{record_id}", response_model=SeaMasterResponse)
async def update_sea_master(
    record_id: UUID,
    request: SeaMasterUpdateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Update a Sea Master record."""
    result = await session.execute(
        select(SeaMasterModel).where(SeaMasterModel.id == str(record_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Sea Master not found")
    if request.to_country_id is not None:
        record.to_country_id = str(request.to_country_id)
    if request.product_type is not None:
        record.product_type = request.product_type
    if request.slab_name is not None:
        record.slab_name = request.slab_name
    if request.currency is not None:
        record.currency = request.currency
    if request.is_active is not None:
        record.is_active = request.is_active
    record.modified_by = current_user.username
    await session.flush()
    res = await session.execute(
        select(SeaMasterModel)
        .options(selectinload(SeaMasterModel.to_country), selectinload(SeaMasterModel.rates))
        .where(SeaMasterModel.id == str(record_id))
    )
    return _sea_to_resp(res.scalar_one())


@router.delete("/sea-masters/{record_id}", status_code=status.HTTP_200_OK)
async def delete_sea_master(
    record_id: UUID,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Soft-delete a Sea Master record."""
    result = await session.execute(
        select(SeaMasterModel).where(SeaMasterModel.id == str(record_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Sea Master not found")
    record.is_active = False
    record.modified_by = current_user.username
    await session.flush()
    return {"message": "Sea Master deactivated successfully"}


# --- Sea Master Rate sub-endpoints ---

@router.post("/sea-masters/{master_id}/rates", response_model=RateResponse, status_code=status.HTTP_201_CREATED)
async def add_sea_master_rate(
    master_id: UUID,
    request: RateCreateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Add a rate to a Sea Master record."""
    check = await session.execute(
        select(SeaMasterModel).where(SeaMasterModel.id == str(master_id))
    )
    if not check.scalar_one_or_none():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Sea Master not found")
    vf = datetime.fromisoformat(request.valid_from.replace("Z", "+00:00"))
    vt = datetime.fromisoformat(request.valid_till.replace("Z", "+00:00"))
    record = SeaMasterRateModel(
        sea_master_id=str(master_id),
        rate=Decimal(str(request.rate)),
        valid_from=vf,
        valid_till=vt,
        is_active=True,
        created_by=current_user.username,
        modified_by=current_user.username,
    )
    session.add(record)
    await session.flush()
    return RateResponse(
        id=record.id,
        rate=float(record.rate),
        valid_from=record.valid_from.isoformat(),
        valid_till=record.valid_till.isoformat(),
        is_active=record.is_active,
    )


@router.delete("/sea-masters/rates/{rate_id}", status_code=status.HTTP_200_OK)
async def delete_sea_master_rate(
    rate_id: UUID,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Soft-delete a Sea Master rate."""
    result = await session.execute(
        select(SeaMasterRateModel).where(SeaMasterRateModel.id == str(rate_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Rate not found")
    record.is_active = False
    record.modified_by = current_user.username
    await session.flush()
    return {"message": "Rate deactivated successfully"}


# ============================================================
# SEA CBM MASTER — CRUD
# ============================================================

@router.get("/sea-cbm", response_model=list[SeaCbmResponse])
async def list_sea_cbm(
    include_inactive: bool = False,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """List Sea CBM Masters."""
    query = select(SeaCbmMasterModel).order_by(SeaCbmMasterModel.from_cbm)
    if not include_inactive:
        query = query.where(SeaCbmMasterModel.is_active == True)  # noqa: E712
    result = await session.execute(query)
    return [SeaCbmResponse.model_validate(r) for r in result.scalars().all()]


@router.post("/sea-cbm", response_model=SeaCbmResponse, status_code=status.HTTP_201_CREATED)
async def create_sea_cbm(
    request: SeaCbmCreateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Create a new Sea CBM Master record."""
    record = SeaCbmMasterModel(
        slab_name=request.slab_name,
        from_cbm=request.from_cbm,
        to_cbm=request.to_cbm,
        type=request.type,
        total_count=request.total_count,
        is_active=True,
        created_by=current_user.username,
        modified_by=current_user.username,
    )
    session.add(record)
    await session.flush()
    await session.refresh(record)
    return SeaCbmResponse.model_validate(record)


@router.put("/sea-cbm/{record_id}", response_model=SeaCbmResponse)
async def update_sea_cbm(
    record_id: UUID,
    request: SeaCbmUpdateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Update a Sea CBM Master record."""
    result = await session.execute(
        select(SeaCbmMasterModel).where(SeaCbmMasterModel.id == str(record_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Sea CBM not found")
    for field, value in request.model_dump(exclude_unset=True).items():
        if value is not None:
            setattr(record, field, value)
    record.modified_by = current_user.username
    await session.flush()
    await session.refresh(record)
    return SeaCbmResponse.model_validate(record)


@router.delete("/sea-cbm/{record_id}", status_code=status.HTTP_200_OK)
async def delete_sea_cbm(
    record_id: UUID,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Soft-delete a Sea CBM Master record."""
    result = await session.execute(
        select(SeaCbmMasterModel).where(SeaCbmMasterModel.id == str(record_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Sea CBM not found")
    record.is_active = False
    record.modified_by = current_user.username
    await session.flush()
    return {"message": "Sea CBM deactivated successfully"}


# ============================================================
# VEHICLE TYPE MASTER — CRUD
# ============================================================

@router.get("/vehicle-types", response_model=list[VehicleTypeResponse])
async def list_vehicle_types(
    include_inactive: bool = False,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """List Vehicle Types."""
    query = select(VehicleTypeMasterModel).order_by(VehicleTypeMasterModel.vehicle_type)
    if not include_inactive:
        query = query.where(VehicleTypeMasterModel.is_active == True)  # noqa: E712
    result = await session.execute(query)
    return [VehicleTypeResponse.model_validate(r) for r in result.scalars().all()]


@router.post("/vehicle-types", response_model=VehicleTypeResponse, status_code=status.HTTP_201_CREATED)
async def create_vehicle_type(
    request: VehicleTypeCreateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Create a new Vehicle Type record."""
    record = VehicleTypeMasterModel(
        from_no_of_pallet_or_box=request.from_no_of_pallet_or_box,
        to_no_of_pallet_or_box=request.to_no_of_pallet_or_box,
        vehicle_type=request.vehicle_type,
        is_active=True,
        created_by=current_user.username,
        modified_by=current_user.username,
    )
    session.add(record)
    await session.flush()
    await session.refresh(record)
    return VehicleTypeResponse.model_validate(record)


@router.put("/vehicle-types/{record_id}", response_model=VehicleTypeResponse)
async def update_vehicle_type(
    record_id: UUID,
    request: VehicleTypeUpdateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Update a Vehicle Type record."""
    result = await session.execute(
        select(VehicleTypeMasterModel).where(VehicleTypeMasterModel.id == str(record_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Vehicle Type not found")
    for field, value in request.model_dump(exclude_unset=True).items():
        if value is not None:
            setattr(record, field, value)
    record.modified_by = current_user.username
    await session.flush()
    await session.refresh(record)
    return VehicleTypeResponse.model_validate(record)


@router.delete("/vehicle-types/{record_id}", status_code=status.HTTP_200_OK)
async def delete_vehicle_type(
    record_id: UUID,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Soft-delete a Vehicle Type record."""
    result = await session.execute(
        select(VehicleTypeMasterModel).where(VehicleTypeMasterModel.id == str(record_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Vehicle Type not found")
    record.is_active = False
    record.modified_by = current_user.username
    await session.flush()
    return {"message": "Vehicle Type deactivated successfully"}


# ============================================================
# LOCAL MASTER — CRUD
# ============================================================

@router.get("/local-masters", response_model=list[LocalMasterResponse])
async def list_local_masters(
    include_inactive: bool = False,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """List Local Masters with rates."""
    query = (
        select(LocalMasterModel)
        .options(
            selectinload(LocalMasterModel.from_city),
            selectinload(LocalMasterModel.to_city),
            selectinload(LocalMasterModel.vehicle_type_ref),
            selectinload(LocalMasterModel.rates),
        )
        .order_by(LocalMasterModel.created_date.desc())
    )
    if not include_inactive:
        query = query.where(LocalMasterModel.is_active == True)  # noqa: E712
    result = await session.execute(query)
    return [_local_to_resp(r) for r in result.scalars().all()]


@router.post("/local-masters", response_model=LocalMasterResponse, status_code=status.HTTP_201_CREATED)
async def create_local_master(
    request: LocalMasterCreateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Create a new Local Master record."""
    record = LocalMasterModel(
        from_city_id=str(request.from_city_id),
        to_city_id=str(request.to_city_id),
        product_type=request.product_type,
        vehicle_type_id=str(request.vehicle_type_id) if request.vehicle_type_id else None,
        min_slab=request.min_slab,
        max_slab=request.max_slab,
        is_active=True,
        created_by=current_user.username,
        modified_by=current_user.username,
    )
    session.add(record)
    await session.flush()
    res = await session.execute(
        select(LocalMasterModel)
        .options(
            selectinload(LocalMasterModel.from_city),
            selectinload(LocalMasterModel.to_city),
            selectinload(LocalMasterModel.vehicle_type_ref),
            selectinload(LocalMasterModel.rates),
        )
        .where(LocalMasterModel.id == record.id)
    )
    return _local_to_resp(res.scalar_one())


@router.put("/local-masters/{record_id}", response_model=LocalMasterResponse)
async def update_local_master(
    record_id: UUID,
    request: LocalMasterUpdateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Update a Local Master record."""
    result = await session.execute(
        select(LocalMasterModel).where(LocalMasterModel.id == str(record_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Local Master not found")
    if request.from_city_id is not None:
        record.from_city_id = str(request.from_city_id)
    if request.to_city_id is not None:
        record.to_city_id = str(request.to_city_id)
    if request.product_type is not None:
        record.product_type = request.product_type
    if request.vehicle_type_id is not None:
        record.vehicle_type_id = str(request.vehicle_type_id)
    if request.min_slab is not None:
        record.min_slab = request.min_slab
    if request.max_slab is not None:
        record.max_slab = request.max_slab
    if request.is_active is not None:
        record.is_active = request.is_active
    record.modified_by = current_user.username
    await session.flush()
    res = await session.execute(
        select(LocalMasterModel)
        .options(
            selectinload(LocalMasterModel.from_city),
            selectinload(LocalMasterModel.to_city),
            selectinload(LocalMasterModel.vehicle_type_ref),
            selectinload(LocalMasterModel.rates),
        )
        .where(LocalMasterModel.id == str(record_id))
    )
    return _local_to_resp(res.scalar_one())


@router.delete("/local-masters/{record_id}", status_code=status.HTTP_200_OK)
async def delete_local_master(
    record_id: UUID,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Soft-delete a Local Master record."""
    result = await session.execute(
        select(LocalMasterModel).where(LocalMasterModel.id == str(record_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Local Master not found")
    record.is_active = False
    record.modified_by = current_user.username
    await session.flush()
    return {"message": "Local Master deactivated successfully"}


# --- Local Master Rate sub-endpoints ---

@router.post("/local-masters/{master_id}/rates", response_model=RateResponse, status_code=status.HTTP_201_CREATED)
async def add_local_master_rate(
    master_id: UUID,
    request: RateCreateRequest,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Add a rate to a Local Master record."""
    check = await session.execute(
        select(LocalMasterModel).where(LocalMasterModel.id == str(master_id))
    )
    if not check.scalar_one_or_none():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Local Master not found")
    vf = datetime.fromisoformat(request.valid_from.replace("Z", "+00:00"))
    vt = datetime.fromisoformat(request.valid_till.replace("Z", "+00:00"))
    record = LocalMasterRateModel(
        local_master_id=str(master_id),
        rate=Decimal(str(request.rate)),
        valid_from=vf,
        valid_till=vt,
        is_active=True,
        created_by=current_user.username,
        modified_by=current_user.username,
    )
    session.add(record)
    await session.flush()
    return RateResponse(
        id=record.id,
        rate=float(record.rate),
        valid_from=record.valid_from.isoformat(),
        valid_till=record.valid_till.isoformat(),
        is_active=record.is_active,
    )


@router.delete("/local-masters/rates/{rate_id}", status_code=status.HTTP_200_OK)
async def delete_local_master_rate(
    rate_id: UUID,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Soft-delete a Local Master rate."""
    result = await session.execute(
        select(LocalMasterRateModel).where(LocalMasterRateModel.id == str(rate_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Rate not found")
    record.is_active = False
    record.modified_by = current_user.username
    await session.flush()
    return {"message": "Rate deactivated successfully"}


# ============================================================
# EXCEL IMPORT — AIR MASTER
# ============================================================

@router.get("/air-master-template")
async def download_air_master_template(
    current_user: User = Depends(get_current_active_user),
):
    """Download Excel template for Air Master import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AirMaster"
    ws.append(["Country Name", "Country Code", "Product Type", "Slab", "Rate", "Currency", "Min Slab", "Max Slab", "Unit"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=air_master_template.xlsx"},
    )


@router.post("/air-master-import", status_code=status.HTTP_200_OK)
async def import_air_master_excel(
    file: UploadFile = File(...),
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Import Air Master from Excel."""
    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
    ws = wb.active
    created = 0
    skipped = 0

    # Preload country_code -> country_id AND country_name -> country_id
    country_result = await session.execute(
        _text("SELECT id, country_code, country_name FROM country_master WHERE is_active = true")
    )
    country_code_map: dict[str, str] = {}
    country_name_map: dict[str, str] = {}
    for r in country_result.fetchall():
        if r[1]:
            country_code_map[r[1].lower().strip()] = str(r[0])
        if r[2]:
            country_name_map[r[2].lower().strip()] = str(r[0])

    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row:
            continue
        country_name = str(row[0]).strip() if len(row) > 0 and row[0] else ""
        country_code = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        product_type = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        slab_name = str(row[3]).strip() if len(row) > 3 and row[3] else ""
        rate = row[4] if len(row) > 4 and row[4] else None
        min_slab = row[6] if len(row) > 6 and row[6] else None
        max_slab = row[7] if len(row) > 7 and row[7] else None

        if not country_code and not country_name:
            continue

        # Look up country by code first, then fall back to name
        country_id = None
        if country_code:
            country_id = country_code_map.get(country_code.lower().strip())
        if not country_id and country_name:
            country_id = country_name_map.get(country_name.lower().strip())

        if not country_id:
            skipped += 1
            continue

        min_slab_val = int(float(min_slab)) if min_slab else 0
        max_slab_val = int(float(max_slab)) if max_slab else 999999
        slab_label = slab_name or f"{min_slab_val}-{max_slab_val} KG"

        record = AirMasterModel(
            to_country_id=country_id,
            product_type=product_type or None,
            min_slab=min_slab_val,
            max_slab=max_slab_val,
            slab_name=slab_label,
            is_active=True,
            created_by=current_user.username,
            modified_by=current_user.username,
        )
        session.add(record)
        await session.flush()

        if rate:
            now = datetime.now(timezone.utc)
            rate_record = AirMasterRateModel(
                air_master_id=record.id,
                rate=Decimal(str(float(rate))),
                valid_from=now,
                valid_till=datetime(2099, 12, 31, tzinfo=timezone.utc),
                is_active=True,
                created_by=current_user.username,
                modified_by=current_user.username,
            )
            session.add(rate_record)
        created += 1

    await session.flush()
    wb.close()
    return {"message": f"Air Master import complete. Created: {created}, Skipped: {skipped}"}


# ============================================================
# EXCEL IMPORT — SEA MASTER
# ============================================================

@router.get("/sea-master-template")
async def download_sea_master_template(
    current_user: User = Depends(get_current_active_user),
):
    """Download Excel template for Sea Master import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SeaMaster"
    ws.append(["Country Name", "Country Code", "Product Type", "Slab", "Rate", "Currency"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=sea_master_template.xlsx"},
    )


@router.post("/sea-master-import", status_code=status.HTTP_200_OK)
async def import_sea_master_excel(
    file: UploadFile = File(...),
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Import Sea Master from Excel."""
    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
    ws = wb.active
    created = 0
    skipped = 0

    # Preload country_code -> country_id AND country_name -> country_id
    country_result = await session.execute(
        _text("SELECT id, country_code, country_name FROM country_master WHERE is_active = true")
    )
    country_code_map: dict[str, str] = {}
    country_name_map: dict[str, str] = {}
    for r in country_result.fetchall():
        if r[1]:
            country_code_map[r[1].lower().strip()] = str(r[0])
        if r[2]:
            country_name_map[r[2].lower().strip()] = str(r[0])

    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row:
            continue
        country_name = str(row[0]).strip() if len(row) > 0 and row[0] else ""
        country_code = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        product_type = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        slab_name = str(row[3]).strip() if len(row) > 3 and row[3] else ""
        rate = row[4] if len(row) > 4 and row[4] else None
        currency = str(row[5]).strip() if len(row) > 5 and row[5] else ""

        if not country_code and not country_name:
            skipped += 1
            continue

        # Look up country by code first, then fall back to name
        country_id = None
        if country_code:
            country_id = country_code_map.get(country_code.lower().strip())
        if not country_id and country_name:
            country_id = country_name_map.get(country_name.lower().strip())

        if not country_id:
            skipped += 1
            continue

        record = SeaMasterModel(
            to_country_id=country_id,
            product_type=product_type or None,
            slab_name=slab_name or None,
            currency=currency or None,
            is_active=True,
            created_by=current_user.username,
            modified_by=current_user.username,
        )
        session.add(record)
        await session.flush()

        if rate:
            now = datetime.now(timezone.utc)
            rate_record = SeaMasterRateModel(
                sea_master_id=record.id,
                rate=Decimal(str(float(rate))),
                valid_from=now,
                valid_till=datetime(2099, 12, 31, tzinfo=timezone.utc),
                is_active=True,
                created_by=current_user.username,
                modified_by=current_user.username,
            )
            session.add(rate_record)
        created += 1

    await session.flush()
    wb.close()
    return {"message": f"Sea Master import complete. Created: {created}, Skipped: {skipped}"}


# ============================================================
# EXCEL IMPORT — SEA CBM MASTER
# ============================================================

@router.get("/sea-cbm-template")
async def download_sea_cbm_template(
    current_user: User = Depends(get_current_active_user),
):
    """Download Excel template for Sea CBM Master import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SeaCbmMaster"
    ws.append(["Sr No", "Slab Name", "CBM From", "CBM To", "Type", "Total Count"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=sea_cbm_master_template.xlsx"},
    )


@router.post("/sea-cbm-import", status_code=status.HTTP_200_OK)
async def import_sea_cbm_excel(
    file: UploadFile = File(...),
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Import Sea CBM Master from Excel."""
    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
    ws = wb.active
    created = 0

    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row:
            continue
        slab_name = str(row[1]).strip() if len(row) > 1 and row[1] else None
        from_cbm = row[2] if len(row) > 2 and row[2] is not None else None
        to_cbm = row[3] if len(row) > 3 and row[3] is not None else None
        type_val = str(row[4]).strip() if len(row) > 4 and row[4] else None
        total_count = row[5] if len(row) > 5 and row[5] is not None else None

        if from_cbm is None and to_cbm is None:
            continue

        record = SeaCbmMasterModel(
            slab_name=slab_name,
            from_cbm=int(float(from_cbm)) if from_cbm is not None else 0,
            to_cbm=int(float(to_cbm)) if to_cbm is not None else 0,
            type=type_val,
            total_count=int(float(total_count)) if total_count is not None else None,
            is_active=True,
            created_by=current_user.username,
            modified_by=current_user.username,
        )
        session.add(record)
        created += 1

    await session.flush()
    wb.close()
    return {"message": f"Sea CBM Master import complete. Created: {created}"}


# ============================================================
# EXCEL IMPORT — LOCAL MASTER
# ============================================================

@router.get("/local-master-template")
async def download_local_master_template(
    current_user: User = Depends(get_current_active_user),
):
    """Download Excel template for Local Master import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LocalMaster"
    ws.append(["From City", "To City", "Product Type", "Vehicle Type", "Min Slab", "Max Slab"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=local_master_template.xlsx"},
    )


@router.post("/local-master-import", status_code=status.HTTP_200_OK)
async def import_local_master_excel(
    file: UploadFile = File(...),
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
):
    """Import Local Master from Excel."""
    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
    ws = wb.active
    created = 0
    skipped = 0

    # Preload city_name -> city_id (case-insensitive)
    city_result = await session.execute(
        _text("SELECT id, city_name FROM city_master WHERE is_active = true")
    )
    city_map = {r[1].lower().strip(): str(r[0]) for r in city_result.fetchall() if r[1]}

    # Preload vehicle_type -> vehicle_type_id (case-insensitive)
    vt_result = await session.execute(
        _text("SELECT id, vehicle_type FROM vehicle_type_master WHERE is_active = true")
    )
    vt_map = {r[1].lower().strip(): str(r[0]) for r in vt_result.fetchall() if r[1]}

    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row:
            continue
        from_city_name = str(row[0]).strip() if len(row) > 0 and row[0] else ""
        to_city_name = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        product_type = str(row[2]).strip() if len(row) > 2 and row[2] else None
        vehicle_type_name = str(row[3]).strip() if len(row) > 3 and row[3] else ""
        min_slab = str(row[4]).strip() if len(row) > 4 and row[4] else None
        max_slab = str(row[5]).strip() if len(row) > 5 and row[5] else None

        if not from_city_name and not to_city_name:
            continue

        from_city_id = city_map.get(from_city_name.lower()) if from_city_name else None
        to_city_id = city_map.get(to_city_name.lower()) if to_city_name else None
        vehicle_type_id = vt_map.get(vehicle_type_name.lower()) if vehicle_type_name else None

        if not from_city_id or not to_city_id:
            skipped += 1
            continue

        # Clean trailing .0 from slab values
        if min_slab and min_slab.endswith(".0"):
            min_slab = min_slab[:-2]
        if max_slab and max_slab.endswith(".0"):
            max_slab = max_slab[:-2]

        record = LocalMasterModel(
            from_city_id=from_city_id,
            to_city_id=to_city_id,
            product_type=product_type,
            vehicle_type_id=vehicle_type_id,
            min_slab=min_slab,
            max_slab=max_slab,
            is_active=True,
            created_by=current_user.username,
            modified_by=current_user.username,
        )
        session.add(record)
        created += 1

    await session.flush()
    wb.close()
    return {"message": f"Local Master import complete. Created: {created}, Skipped: {skipped}"}
