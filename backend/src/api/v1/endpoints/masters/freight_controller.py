"""
Freight Master Data CRUD + Excel Import endpoints.
Provides Currency Conversion, Container, Transit Day, and Fixed Charges management.
"""

from decimal import Decimal
from io import BytesIO
from uuid import UUID

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from src.api.v1.dependencies import get_current_active_user
from src.domain.entities.user import User
from src.infrastructure.database.session import get_db_session
from src.infrastructure.database.models.masters.currency_conversion_model import (
    CurrencyConversionRateModel,
)
from src.infrastructure.database.models.masters.container_model import (
    ContainerMasterModel,
)
from src.infrastructure.database.models.masters.transit_day_model import (
    TransitDayMasterModel,
)
from src.infrastructure.database.models.masters.fixed_charges_model import (
    FixedChargesMasterModel,
)
from src.api.v1.endpoints.masters.freight_schemas import (
    CurrencyConversionCreate,
    CurrencyConversionUpdate,
    CurrencyConversionResponse,
    ContainerCreate,
    ContainerUpdate,
    ContainerResponse,
    TransitDayCreate,
    TransitDayUpdate,
    TransitDayResponse,
    FixedChargesCreate,
    FixedChargesUpdate,
    FixedChargesResponse,
)

router = APIRouter(prefix="/masters", tags=["Freight Master Data"])


# ============================================================
# CURRENCY CONVERSION — CRUD
# ============================================================


@router.get("/currency-conversions", response_model=list[CurrencyConversionResponse])
async def list_currency_conversions(
    include_inactive: bool = False,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> list[CurrencyConversionResponse]:
    """List all currency conversion rates."""
    query = select(CurrencyConversionRateModel).order_by(
        CurrencyConversionRateModel.created_date.desc()
    )
    if not include_inactive:
        query = query.where(CurrencyConversionRateModel.is_active == True)  # noqa: E712
    result = await session.execute(query)
    items = []
    for r in result.scalars().all():
        items.append(CurrencyConversionResponse(
            id=r.id,
            exrt=r.exrt,
            from_currency=r.from_currency,
            to_currency=r.to_currency,
            valid_from=r.valid_from.isoformat() if r.valid_from else None,
            exchange_rate=float(r.exchange_rate) if r.exchange_rate else None,
            ratio_from=r.ratio_from,
            ratio_to=r.ratio_to,
            is_active=r.is_active,
        ))
    return items


@router.post(
    "/currency-conversions",
    response_model=CurrencyConversionResponse,
    status_code=status.HTTP_201_CREATED,
)
async def create_currency_conversion(
    request: CurrencyConversionCreate,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> CurrencyConversionResponse:
    """Create a new currency conversion rate."""
    from datetime import datetime as _dt

    vf = (
        _dt.fromisoformat(request.valid_from.replace("Z", "+00:00"))
        if request.valid_from
        else None
    )
    record = CurrencyConversionRateModel(
        exrt=request.exrt,
        from_currency=request.from_currency,
        to_currency=request.to_currency,
        valid_from=vf,
        exchange_rate=Decimal(str(request.exchange_rate)) if request.exchange_rate else None,
        ratio_from=request.ratio_from,
        ratio_to=request.ratio_to,
        is_active=True,
        created_by=current_user.username,
        modified_by=current_user.username,
    )
    session.add(record)
    await session.flush()
    return CurrencyConversionResponse(
        id=record.id,
        exrt=record.exrt,
        from_currency=record.from_currency,
        to_currency=record.to_currency,
        valid_from=record.valid_from.isoformat() if record.valid_from else None,
        exchange_rate=float(record.exchange_rate) if record.exchange_rate else None,
        ratio_from=record.ratio_from,
        ratio_to=record.ratio_to,
        is_active=record.is_active,
    )


@router.put("/currency-conversions/{rid}", response_model=CurrencyConversionResponse)
async def update_currency_conversion(
    rid: UUID,
    request: CurrencyConversionUpdate,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> CurrencyConversionResponse:
    """Update an existing currency conversion rate."""
    from datetime import datetime as _dt

    result = await session.execute(
        select(CurrencyConversionRateModel).where(CurrencyConversionRateModel.id == rid)
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Not found")
    if request.exrt is not None:
        record.exrt = request.exrt
    if request.from_currency is not None:
        record.from_currency = request.from_currency
    if request.to_currency is not None:
        record.to_currency = request.to_currency
    if request.valid_from is not None:
        record.valid_from = _dt.fromisoformat(request.valid_from.replace("Z", "+00:00"))
    if request.exchange_rate is not None:
        record.exchange_rate = Decimal(str(request.exchange_rate))
    if request.ratio_from is not None:
        record.ratio_from = request.ratio_from
    if request.ratio_to is not None:
        record.ratio_to = request.ratio_to
    if request.is_active is not None:
        record.is_active = request.is_active
    record.modified_by = current_user.username
    await session.flush()
    return CurrencyConversionResponse(
        id=record.id,
        exrt=record.exrt,
        from_currency=record.from_currency,
        to_currency=record.to_currency,
        valid_from=record.valid_from.isoformat() if record.valid_from else None,
        exchange_rate=float(record.exchange_rate) if record.exchange_rate else None,
        ratio_from=record.ratio_from,
        ratio_to=record.ratio_to,
        is_active=record.is_active,
    )


@router.delete("/currency-conversions/{rid}", status_code=200)
async def delete_currency_conversion(
    rid: UUID,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> dict:
    """Soft-delete (deactivate) a currency conversion rate."""
    result = await session.execute(
        select(CurrencyConversionRateModel).where(CurrencyConversionRateModel.id == rid)
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Not found")
    record.is_active = False
    record.modified_by = current_user.username
    await session.flush()
    return {"message": "Deactivated successfully"}


@router.post("/currency-conversions/import-sap")
async def import_currency_conversions_from_sap(
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> dict:
    """Import exchange rates from SAP Integration Suite."""
    import httpx
    from datetime import datetime as _dt
    from sqlalchemy import text as _text

    SAP_URL = "https://integration-suite-test-874gff1p.it-cpi021-rt.cfapps.in30.hana.ondemand.com/http/QAS/Get/ExchangeRates"
    SAP_USER = "sb-03d382c0-4ab1-4efa-b432-d35e9c5c4882!b4582|it-rt-integration-suite-test-874gff1p!b148"
    SAP_PASS = "c4578406-ada9-4508-a1b0-678f77e4a153$qGuIICPVCIlAXVJiXauliATx_mJPAMlPZgzDgoM9diQ="

    try:
        async with httpx.AsyncClient(timeout=30.0, verify=False) as client:
            response = await client.get(SAP_URL, auth=(SAP_USER, SAP_PASS))
            response.raise_for_status()
            data = response.json()
    except httpx.HTTPStatusError as e:
        raise HTTPException(
            status_code=502, detail=f"SAP returned error: {e.response.text[:200]}"
        )
    except Exception as e:
        raise HTTPException(
            status_code=502, detail=f"Failed to fetch from SAP: {str(e)}"
        )

    items = data.get("IT_RESULT", {}).get("item", [])
    if not items:
        return {"message": "No exchange rate records returned from SAP"}

    # Deactivate all existing rates before importing fresh
    await session.execute(
        _text("UPDATE currency_conversion_rate_master SET is_active = false")
    )

    created = 0
    for item in items:
        from_currency = item.get("FCURR") or ""
        to_currency = item.get("TCURR") or ""
        if not from_currency or not to_currency:
            continue

        exrt = item.get("KURST") or ""
        gdatu = item.get("GDATU") or ""
        ukurs = item.get("UKURS") or "0"

        # Parse date YYYYMMDD
        valid_from = None
        if gdatu and len(gdatu) == 8:
            try:
                valid_from = _dt(int(gdatu[:4]), int(gdatu[4:6]), int(gdatu[6:8]))
            except Exception:
                pass

        exchange_rate = Decimal(str(float(ukurs))) if ukurs else Decimal(0)

        session.add(CurrencyConversionRateModel(
            exrt=exrt or None,
            from_currency=from_currency,
            to_currency=to_currency,
            valid_from=valid_from,
            exchange_rate=exchange_rate,
            is_active=True,
            created_by=current_user.username,
            modified_by=current_user.username,
        ))
        created += 1

    await session.flush()
    return {"message": f"SAP Import complete. Imported: {created} exchange rates"}


# ============================================================
# CONTAINER MASTER — CRUD
# ============================================================


@router.get("/containers", response_model=list[ContainerResponse])
async def list_containers(
    include_inactive: bool = False,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> list[ContainerResponse]:
    """List all containers."""
    query = select(ContainerMasterModel).order_by(ContainerMasterModel.container_type)
    if not include_inactive:
        query = query.where(ContainerMasterModel.is_active == True)  # noqa: E712
    result = await session.execute(query)
    return [ContainerResponse.model_validate(r) for r in result.scalars().all()]


@router.post("/containers", response_model=ContainerResponse, status_code=status.HTTP_201_CREATED)
async def create_container(
    request: ContainerCreate,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> ContainerResponse:
    """Create a new container type."""
    existing = await session.execute(
        select(ContainerMasterModel).where(
            ContainerMasterModel.container_type.ilike(request.container_type)
        )
    )
    if existing.scalar_one_or_none():
        raise HTTPException(status_code=400, detail="Container type already exists.")
    record = ContainerMasterModel(
        container_type=request.container_type,
        is_active=True,
        created_by=current_user.username,
        modified_by=current_user.username,
    )
    session.add(record)
    await session.flush()
    return ContainerResponse.model_validate(record)


@router.put("/containers/{rid}", response_model=ContainerResponse)
async def update_container(
    rid: UUID,
    request: ContainerUpdate,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> ContainerResponse:
    """Update an existing container type."""
    result = await session.execute(
        select(ContainerMasterModel).where(ContainerMasterModel.id == rid)
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Not found")
    if request.container_type is not None:
        record.container_type = request.container_type
    if request.is_active is not None:
        record.is_active = request.is_active
    record.modified_by = current_user.username
    await session.flush()
    return ContainerResponse.model_validate(record)


@router.delete("/containers/{rid}", status_code=200)
async def delete_container(
    rid: UUID,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> dict:
    """Soft-delete (deactivate) a container type."""
    result = await session.execute(
        select(ContainerMasterModel).where(ContainerMasterModel.id == rid)
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Not found")
    record.is_active = False
    record.modified_by = current_user.username
    await session.flush()
    return {"message": "Container deactivated successfully"}


# ============================================================
# TRANSIT DAY MASTER — CRUD + IMPORT
# ============================================================


@router.get("/transit-days", response_model=list[TransitDayResponse])
async def list_transit_days(
    include_inactive: bool = False,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> list[TransitDayResponse]:
    """List all transit days."""
    query = select(TransitDayMasterModel).order_by(TransitDayMasterModel.country_name)
    if not include_inactive:
        query = query.where(TransitDayMasterModel.is_active == True)  # noqa: E712
    result = await session.execute(query)
    return [TransitDayResponse.model_validate(r) for r in result.scalars().all()]


@router.post("/transit-days", response_model=TransitDayResponse, status_code=status.HTTP_201_CREATED)
async def create_transit_day(
    request: TransitDayCreate,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> TransitDayResponse:
    """Create a new transit day record."""
    record = TransitDayMasterModel(
        country_name=request.country_name,
        country_code=request.country_code,
        customer_clearance=request.customer_clearance,
        transit_days_for_air=request.transit_days_for_air,
        container_load_for_sea=request.container_load_for_sea,
        test_days=request.test_days,
        is_active=True,
        created_by=current_user.username,
        modified_by=current_user.username,
    )
    session.add(record)
    await session.flush()
    return TransitDayResponse.model_validate(record)


@router.put("/transit-days/{rid}", response_model=TransitDayResponse)
async def update_transit_day(
    rid: UUID,
    request: TransitDayUpdate,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> TransitDayResponse:
    """Update an existing transit day record."""
    result = await session.execute(
        select(TransitDayMasterModel).where(TransitDayMasterModel.id == rid)
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Not found")
    for field, value in request.model_dump(exclude_unset=True).items():
        if value is not None:
            setattr(record, field, value)
    record.modified_by = current_user.username
    await session.flush()
    return TransitDayResponse.model_validate(record)


@router.delete("/transit-days/{rid}", status_code=200)
async def delete_transit_day(
    rid: UUID,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> dict:
    """Soft-delete (deactivate) a transit day record."""
    result = await session.execute(
        select(TransitDayMasterModel).where(TransitDayMasterModel.id == rid)
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Not found")
    record.is_active = False
    record.modified_by = current_user.username
    await session.flush()
    return {"message": "Transit Day deactivated successfully"}


@router.get("/transit-days-template")
async def download_transit_days_template() -> StreamingResponse:
    """Download Excel template for transit days import."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TransitDays"
    ws.append(["Country Name", "Country Code", "Customer Clearance", "Transit Days (Air)", "Transit Days (Sea)", "Test Days"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=transit_days_template.xlsx"},
    )


@router.post("/transit-days-import", status_code=status.HTTP_200_OK)
async def import_transit_days_excel(
    file: UploadFile = File(...),
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> dict:
    """Import Transit Days from Excel. Columns: A=Country Name, B=Country Code, C=Customer Clearance, D=Transit Days Air, E=Transit Days Sea, F=Test Days."""
    import openpyxl

    contents = await file.read()
    wb = openpyxl.load_workbook(BytesIO(contents), read_only=True)
    ws = wb.active
    created = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        country_name = str(row[0]).strip() if row[0] else None
        if not country_name:
            continue
        country_code = str(row[1]).strip() if len(row) > 1 and row[1] else None
        customer_clearance = int(float(row[2])) if len(row) > 2 and row[2] is not None else None
        transit_days_air = int(float(row[3])) if len(row) > 3 and row[3] is not None else None
        transit_days_sea = int(float(row[4])) if len(row) > 4 and row[4] is not None else None
        test_days = int(float(row[5])) if len(row) > 5 and row[5] is not None else None

        session.add(TransitDayMasterModel(
            country_name=country_name,
            country_code=country_code,
            customer_clearance=customer_clearance,
            transit_days_for_air=transit_days_air,
            container_load_for_sea=transit_days_sea,
            test_days=test_days,
            is_active=True,
            created_by=current_user.username,
            modified_by=current_user.username,
        ))
        created += 1

    await session.flush()
    wb.close()
    return {"message": f"Transit Days import complete. Created: {created}"}


# ============================================================
# FIXED CHARGES MASTER — CRUD + IMPORT
# ============================================================


@router.get("/fixed-charges", response_model=list[FixedChargesResponse])
async def list_fixed_charges(
    include_inactive: bool = False,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> list[FixedChargesResponse]:
    """List all fixed charges."""
    query = select(FixedChargesMasterModel).order_by(
        FixedChargesMasterModel.created_date.desc()
    )
    if not include_inactive:
        query = query.where(FixedChargesMasterModel.is_active == True)  # noqa: E712
    result = await session.execute(query)
    return [FixedChargesResponse.model_validate(r) for r in result.scalars().all()]


@router.post("/fixed-charges", response_model=FixedChargesResponse, status_code=status.HTTP_201_CREATED)
async def create_fixed_charges(
    request: FixedChargesCreate,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> FixedChargesResponse:
    """Create a new fixed charges record."""
    data = request.model_dump()
    for f in ["document_charges", "shrink_wrap_charges", "unloading_charges", "pallet_charges", "data_logger_charges", "blanket_charges"]:
        if data.get(f) is not None:
            data[f] = Decimal(str(data[f]))
    record = FixedChargesMasterModel(
        **data,
        is_active=True,
        created_by=current_user.username,
        modified_by=current_user.username,
    )
    session.add(record)
    await session.flush()
    return FixedChargesResponse.model_validate(record)


@router.put("/fixed-charges/{rid}", response_model=FixedChargesResponse)
async def update_fixed_charges(
    rid: UUID,
    request: FixedChargesUpdate,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> FixedChargesResponse:
    """Update an existing fixed charges record."""
    result = await session.execute(
        select(FixedChargesMasterModel).where(FixedChargesMasterModel.id == rid)
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Not found")
    decimal_fields = {"document_charges", "shrink_wrap_charges", "unloading_charges", "pallet_charges", "data_logger_charges", "blanket_charges"}
    for f, v in request.model_dump(exclude_unset=True).items():
        if v is not None:
            if f in decimal_fields:
                setattr(record, f, Decimal(str(v)))
            else:
                setattr(record, f, v)
    record.modified_by = current_user.username
    await session.flush()
    return FixedChargesResponse.model_validate(record)


@router.delete("/fixed-charges/{rid}", status_code=200)
async def delete_fixed_charges(
    rid: UUID,
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> dict:
    """Soft-delete (deactivate) a fixed charges record."""
    result = await session.execute(
        select(FixedChargesMasterModel).where(FixedChargesMasterModel.id == rid)
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Not found")
    record.is_active = False
    record.modified_by = current_user.username
    await session.flush()
    return {"message": "Fixed Charges deactivated successfully"}


@router.get("/fixed-charges-template")
async def download_fixed_charges_template() -> StreamingResponse:
    """Download Excel template for fixed charges import."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FixedCharges"
    ws.append(["Mode of Shipment", "Pallet Type", "Document Charges", "Shrink Wrap Charges", "Unloading Charges", "Pallet Charges", "Data Logger Charges", "Blanket Charges"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=fixed_charges_template.xlsx"},
    )


@router.post("/fixed-charges-import", status_code=status.HTTP_200_OK)
async def import_fixed_charges_excel(
    file: UploadFile = File(...),
    session: AsyncSession = Depends(get_db_session),
    current_user: User = Depends(get_current_active_user),
) -> dict:
    """Import Fixed Charges from Excel. Columns: A=Mode of Shipment, B=Pallet Type, C=Document, D=Shrink Wrap, E=Unloading, F=Pallet, G=Data Logger, H=Blanket."""
    import openpyxl

    contents = await file.read()
    wb = openpyxl.load_workbook(BytesIO(contents), read_only=True)
    ws = wb.active
    created = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        mode_of_shipment = str(row[0]).strip() if row[0] else None
        if not mode_of_shipment:
            continue
        pallet_name = str(row[1]).strip() if len(row) > 1 and row[1] else None
        document_charges = Decimal(str(float(row[2]))) if len(row) > 2 and row[2] not in (None, "", " ") else None
        shrink_wrap_charges = Decimal(str(float(row[3]))) if len(row) > 3 and row[3] not in (None, "", " ") else None
        unloading_charges = Decimal(str(float(row[4]))) if len(row) > 4 and row[4] not in (None, "", " ") else None
        pallet_charges = Decimal(str(float(row[5]))) if len(row) > 5 and row[5] not in (None, "", " ") else None
        data_logger_charges = Decimal(str(float(row[6]))) if len(row) > 6 and row[6] not in (None, "", " ") else None
        blanket_charges = Decimal(str(float(row[7]))) if len(row) > 7 and row[7] not in (None, "", " ") else None

        session.add(FixedChargesMasterModel(
            mode_of_shipment=mode_of_shipment,
            pallet_name=pallet_name,
            document_charges=document_charges,
            shrink_wrap_charges=shrink_wrap_charges,
            unloading_charges=unloading_charges,
            pallet_charges=pallet_charges,
            data_logger_charges=data_logger_charges,
            blanket_charges=blanket_charges,
            is_active=True,
            created_by=current_user.username,
            modified_by=current_user.username,
        ))
        created += 1

    await session.flush()
    wb.close()
    return {"message": f"Fixed Charges import complete. Created: {created}"}
